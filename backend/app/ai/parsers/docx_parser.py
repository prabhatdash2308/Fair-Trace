import docx
from typing import Dict, Any

from app.ai.parsers.base import BaseParser
from app.ai.parsers.models import ParsedDocument
from app.ai.parsers.exceptions import NoTextFoundError, ParsingFailedError
from app.ai.parsers.text_normalizer import TextNormalizer


class DOCXParser(BaseParser):
    @property
    def parser_name(self) -> str:
        return "python-docx"

    @property
    def parser_version(self) -> str:
        return docx.__version__

    def supports(self, mime_type: str, extension: str) -> bool:
        return mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or extension.lower() == ".docx"

    def parse(self, file_path: str) -> ParsedDocument:
        try:
            doc = docx.Document(file_path)
            
            text_chunks = []
            
            # Headers
            for section in doc.sections:
                for header_para in section.header.paragraphs:
                    if header_para.text.strip():
                        text_chunks.append(header_para.text)

            # Main text
            for para in doc.paragraphs:
                if para.text.strip():
                    text_chunks.append(para.text)
                    
            # Tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text_chunks.append(" | ".join(row_text))

            # Footers
            for section in doc.sections:
                for footer_para in section.footer.paragraphs:
                    if footer_para.text.strip():
                        text_chunks.append(footer_para.text)

            raw_text = "\n\n".join(text_chunks)
            normalized_text = TextNormalizer.normalize(raw_text)
            
            if not normalized_text:
                raise NoTextFoundError("No text could be extracted from the DOCX document.")
                
            metadata = self._extract_metadata(doc)
            
            # Estimate pages (Word docs don't have hard pages in the file structure like PDFs)
            # Typically 250-300 words per page. We'll use 250 for estimation.
            word_count = len(normalized_text.split())
            estimated_pages = max(1, word_count // 250)
            
            return ParsedDocument(
                text=normalized_text,
                page_count=estimated_pages,
                word_count=word_count,
                character_count=len(normalized_text),
                title=metadata.get("title"),
                author=metadata.get("author"),
                metadata=metadata,
                parser_used=self.parser_name,
                parser_version=self.parser_version,
            )
            
        except NoTextFoundError:
            raise
        except Exception as e:
            raise ParsingFailedError(f"Failed to parse DOCX: {str(e)}") from e

    def _extract_metadata(self, doc: docx.Document) -> Dict[str, Any]:
        core_props = doc.core_properties
        return {
            "title": core_props.title if core_props.title else None,
            "author": core_props.author if core_props.author else None,
            "subject": core_props.subject if core_props.subject else None,
            "keywords": core_props.keywords if core_props.keywords else None,
            "category": core_props.category if core_props.category else None,
            "comments": core_props.comments if core_props.comments else None,
        }
