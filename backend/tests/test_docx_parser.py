import pytest
import os
import tempfile
import docx
from app.ai.parsers.docx_parser import DOCXParser
from app.ai.parsers.exceptions import NoTextFoundError

@pytest.fixture
def parser():
    return DOCXParser()

def create_docx(sections: list, suffix=".docx") -> str:
    doc = docx.Document()
    for section_type, content in sections:
        if section_type == "para":
            doc.add_paragraph(content)
        elif section_type == "table":
            table = doc.add_table(rows=len(content), cols=len(content[0]))
            for i, row in enumerate(content):
                for j, cell_text in enumerate(row):
                    table.cell(i, j).text = cell_text
    
    fd, path = tempfile.mkstemp(suffix=suffix)
    os.close(fd)
    doc.save(path)
    return path

def test_parse_docx_paragraphs(parser):
    path = create_docx([("para", "Hello DOCX"), ("para", "Second paragraph")])
    try:
        doc = parser.parse(path)
        assert "Hello DOCX" in doc.text
        assert "Second paragraph" in doc.text
    finally:
        os.remove(path)

def test_parse_docx_tables(parser):
    path = create_docx([("table", [["A1", "B1"], ["A2", "B2"]])])
    try:
        doc = parser.parse(path)
        assert "A1 | B1" in doc.text
        assert "A2 | B2" in doc.text
    finally:
        os.remove(path)

def test_empty_docx_raises_no_text(parser):
    path = create_docx([])
    try:
        with pytest.raises(NoTextFoundError):
            parser.parse(path)
    finally:
        os.remove(path)
